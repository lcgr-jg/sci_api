"""Excel writer for the quirky china_semi_balance desk layout."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="002060")
HEADER_FONT = Font(bold=True, color="FFFFFF")
NUM_FMT = "#,##0.00"
KBD_FMT = '_-* #,##0.00_-;\\-* #,##0.00_-;_-* "-"??_-;_-@_-'
MONTH_FMT = '[$-409]mmm\\-yy;@'

# Hierarchy styling — darker / bolder = higher level.
_LEVEL_STYLE = {
    "section": {
        "fill": PatternFill("solid", fgColor="002060"),
        "font": Font(bold=True, color="FFFFFF", size=11),
        "indent": 0,
    },
    "total": {
        "fill": PatternFill("solid", fgColor="9BC2E6"),
        "font": Font(bold=True, size=10),
        "indent": 0,
        "top_border": True,
    },
    "province": {
        "fill": PatternFill("solid", fgColor="DDEBF7"),
        "font": Font(bold=True, size=10),
        "indent": 0,
    },
    "owner": {
        "fill": PatternFill("solid", fgColor="F2F2F2"),
        "font": Font(bold=True, size=10),
        "indent": 1,
    },
    "detail": {
        "fill": None,
        "font": Font(bold=False, size=10),
        "indent": 2,
    },
}

_THIN_TOP = Border(top=Side(style="thin", color="595959"))


def _style_header(cell) -> None:
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal="center")


def _apply_level_style(
    ws,
    excel_row: int,
    cols: list[int],
    row_level: str,
) -> None:
    style = _LEVEL_STYLE.get(row_level, _LEVEL_STYLE["detail"])
    for col in cols:
        cell = ws.cell(excel_row, col)
        if style["fill"] is not None:
            cell.fill = style["fill"]
        if style.get("top_border"):
            cell.border = _THIN_TOP
        cell.font = style["font"]


def write_semi_balance_workbook(
    rows: pd.DataFrame,
    months: list[pd.Timestamp],
    path: str | Path,
    *,
    barrels_per_mt: float = 7.33,
) -> Path:
    """
    Write desk layout: kt block (left) + kbd block (right) with Excel formulas.

    Expected columns on ``rows``:
      section, sub_category, item, unit, value_kind ('kt'|'kbd'), row_level,
      plus one column per month labelled YYYY-MM-DD.
    ``row_level`` ∈ {section, total, province, owner, detail} controls fill/indent.
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    months = [pd.Timestamp(m).normalize().replace(day=1) for m in months]
    n = len(months)
    if n == 0:
        raise ValueError("No months to write")

    label_cols = {"section": 3, "sub_category": 4, "item": 5, "unit": 6}
    kt_start = 7  # G
    kbd_unit_col = kt_start + n + 4
    kbd_start = kbd_unit_col + 1

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # Row 3: days-in-month (used by kbd formulas).
    ws.cell(3, 6, value=months[0] - pd.offsets.MonthBegin(1))
    ws.cell(3, 6).number_format = MONTH_FMT
    _style_header(ws.cell(3, 6))
    for i, month in enumerate(months):
        cell = ws.cell(3, kt_start + i, value=int(month.days_in_month))
        cell.number_format = "0"

    headers = [
        (label_cols["sub_category"], "Sub-category"),
        (label_cols["item"], "Item"),
        (label_cols["unit"], "Unit"),
    ]
    for col, text in headers:
        cell = ws.cell(4, col, value=text)
        _style_header(cell)

    for i, month in enumerate(months):
        cell = ws.cell(4, kt_start + i, value=month.to_pydatetime())
        cell.number_format = MONTH_FMT
        _style_header(cell)
        cell = ws.cell(4, kbd_start + i, value=month.to_pydatetime())
        cell.number_format = MONTH_FMT
        _style_header(cell)

    month_keys = [m.strftime("%Y-%m-%d") for m in months]
    value_cols = (
        list(range(kt_start, kt_start + n))
        + [kbd_unit_col]
        + list(range(kbd_start, kbd_start + n))
    )
    style_cols = [
        label_cols["section"],
        label_cols["sub_category"],
        label_cols["item"],
        label_cols["unit"],
        *value_cols,
    ]

    excel_row = 5
    prev_section = None
    prev_level = None

    for _, rec in rows.iterrows():
        section = rec.get("section")
        section_txt = (
            None
            if section is None or (isinstance(section, float) and pd.isna(section))
            else str(section)
        )
        row_level = rec.get("row_level") or "detail"
        if isinstance(row_level, float) and pd.isna(row_level):
            row_level = "detail"

        # Blank line between major sections, and between province blocks.
        if section_txt and section_txt != prev_section and prev_section is not None:
            excel_row += 1
        elif (
            row_level == "province"
            and prev_level in {"owner", "detail", "province"}
            and prev_section == section_txt
        ):
            excel_row += 1

        if section_txt and section_txt != prev_section:
            prev_section = section_txt

        sub = rec.get("sub_category")
        item = rec.get("item")
        unit = rec.get("unit")
        kind = rec.get("value_kind", "kt")

        has_item = (
            item is not None
            and not (isinstance(item, float) and pd.isna(item))
            and str(item) != ""
        )
        has_sub = (
            sub is not None
            and not (isinstance(sub, float) and pd.isna(sub))
            and str(sub) != ""
        )

        if section_txt:
            ws.cell(excel_row, label_cols["section"], value=section_txt)
        if has_sub:
            ws.cell(excel_row, label_cols["sub_category"], value=sub)
        if has_item:
            cell = ws.cell(excel_row, label_cols["item"], value=item)
            indent = _LEVEL_STYLE.get(row_level, _LEVEL_STYLE["detail"])["indent"]
            cell.alignment = Alignment(indent=indent)

        # Section banners: label-only row with full-width fill.
        if row_level == "section" or (not has_item and not has_sub):
            _apply_level_style(ws, excel_row, style_cols, "section")
            if section_txt:
                ws.cell(excel_row, label_cols["section"]).font = _LEVEL_STYLE["section"]["font"]
            excel_row += 1
            prev_level = "section"
            continue

        if pd.notna(unit) and kind == "kt":
            ws.cell(excel_row, label_cols["unit"], value=unit)
        elif kind == "kbd" and pd.notna(unit):
            ws.cell(excel_row, label_cols["unit"], value=unit)
        ws.cell(excel_row, kbd_unit_col, value="kbd")

        for i, key in enumerate(month_keys):
            val = rec.get(key)
            if kind == "kbd":
                if val is not None and pd.notna(val):
                    cell = ws.cell(excel_row, kbd_start + i, value=float(val))
                    cell.number_format = KBD_FMT
            else:
                if val is not None and pd.notna(val):
                    cell = ws.cell(excel_row, kt_start + i, value=float(val))
                    cell.number_format = NUM_FMT
                kt_ref = f"{get_column_letter(kt_start + i)}{excel_row}"
                days_ref = f"{get_column_letter(kt_start + i)}$3"
                formula = f"={kt_ref}*{barrels_per_mt}/({days_ref})"
                cell = ws.cell(excel_row, kbd_start + i, value=formula)
                cell.number_format = KBD_FMT

        _apply_level_style(ws, excel_row, style_cols, str(row_level))
        # Re-apply item indent after fill/font (openpyxl alignment is per-cell).
        if has_item:
            indent = _LEVEL_STYLE.get(row_level, _LEVEL_STYLE["detail"])["indent"]
            ws.cell(excel_row, label_cols["item"]).alignment = Alignment(indent=indent)

        prev_level = str(row_level)
        excel_row += 1

    ws.column_dimensions["C"].width = 32
    ws.column_dimensions["D"].width = 36
    ws.column_dimensions["E"].width = 44
    ws.column_dimensions["F"].width = 8
    for i in range(n):
        ws.column_dimensions[get_column_letter(kt_start + i)].width = 12
        ws.column_dimensions[get_column_letter(kbd_start + i)].width = 12
    ws.column_dimensions[get_column_letter(kbd_unit_col)].width = 6

    # Freeze panes: labels + header rows stay visible while scrolling months.
    ws.freeze_panes = "G5"

    wb.save(path)
    return path
