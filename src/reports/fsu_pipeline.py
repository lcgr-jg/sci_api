"""Parse Petro-Logistics FSU Excel for Russia→China pipeline crude (bpd)."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

# Desk labels used in china_semi_balance / FSU Crude Oil sheet.
DEFAULT_CHINA_PIPELINE_ROUTES = (
    "China via Kazakhstan",
    "China via ESPO",
)


def resolve_fsu_path(path_or_glob: str | Path, *, search_dir: Path | None = None) -> Path:
    """Resolve a concrete FSU workbook path (supports a glob under search_dir)."""
    raw = Path(path_or_glob)
    if raw.is_file():
        return raw.resolve()

    base = search_dir or Path.cwd()
    matches = sorted(base.glob(str(path_or_glob)))
    if not matches:
        matches = sorted(base.glob(raw.name))
    files = [p for p in matches if p.is_file()]
    if not files:
        raise FileNotFoundError(
            f"No FSU workbook matched {path_or_glob!r} under {base}"
        )
    # Prefer the lexicographically last match (usually newest YYYY-MM in the name).
    return files[-1].resolve()


def parse_fsu_china_pipeline(
    path: str | Path,
    *,
    sheet_name: str = "Crude Oil",
    routes: tuple[str, ...] = DEFAULT_CHINA_PIPELINE_ROUTES,
) -> pd.DataFrame:
    """
    Read monthly China pipeline crude from an FSU bpd interim workbook.

    Returns long frame: month_start, route, bpd, kbd (= bpd / 1000).
    Uses the Crude Oil sheet historical month columns (row 4 datetimes),
    which match the desk sample values (Russia Exports can differ).
    """
    path = Path(path)
    wb = load_workbook(path, data_only=True, read_only=True)
    if sheet_name not in wb.sheetnames:
        raise KeyError(f"Sheet {sheet_name!r} not in {path.name}: {wb.sheetnames}")
    ws = wb[sheet_name]

    # Row 4 holds month-start datetimes. The left "spotlight" block (latest month,
    # prior month, YoY) duplicates months that also appear in the long history
    # block to the right — keep the rightmost column for each month.
    month_by_period: dict[pd.Timestamp, int] = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(4, col).value
        ts = pd.to_datetime(val, errors="coerce")
        if pd.isna(ts):
            continue
        month = pd.Timestamp(ts).normalize().replace(day=1)
        month_by_period[month] = col  # later columns overwrite earlier duplicates

    month_cols = {col: month for month, col in month_by_period.items()}
    if not month_cols:
        raise ValueError(f"No month headers found on row 4 of {sheet_name!r}")

    wanted = {r.lower(): r for r in routes}
    found: dict[str, int] = {}
    for row in range(1, ws.max_row + 1):
        label = ws.cell(row, 1).value
        if label is None:
            continue
        key = str(label).strip().lower()
        if key in wanted and key not in found:
            found[key] = row

    missing = [wanted[k] for k in wanted if k not in found]
    if missing:
        raise KeyError(f"Routes not found in {sheet_name!r}: {missing}")

    records: list[dict] = []
    for key, row_idx in found.items():
        route = wanted[key]
        for col, month_start in month_cols.items():
            raw = ws.cell(row_idx, col).value
            bpd = pd.to_numeric(raw, errors="coerce")
            if pd.isna(bpd):
                continue
            records.append(
                {
                    "month_start": month_start,
                    "route": route,
                    "bpd": float(bpd),
                    "kbd": float(bpd) / 1000.0,
                }
            )

    wb.close()
    if not records:
        return pd.DataFrame(columns=["month_start", "route", "bpd", "kbd"])

    return (
        pd.DataFrame(records)
        .sort_values(["route", "month_start"])
        .reset_index(drop=True)
    )
